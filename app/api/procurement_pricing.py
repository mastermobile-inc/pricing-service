import csv
from datetime import date
from io import BytesIO, StringIO
from typing import Literal

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.api.dependencies import get_db
from app.models.procurement_pricing import ProcurementPriceBatch, ProcurementPricePreset
from app.schemas.procurement_pricing import (
    PriceApproval,
    PriceBatchCreate,
    PriceBatchRead,
    PriceHistoryPoint,
    PricePresetRead,
    PricePresetWrite,
    PricingFilter,
    PricingTable,
)
from app.services.bitrix_procurement_order_formation_auth import (
    ProcurementOrderFormationSession,
    ensure_bitrix_user_allowed,
    verify_procurement_order_formation_session,
)
from app.services.procurement_pricing import (
    PricingConflict,
    actor_key,
    approve_batch,
    create_batch,
    get_batch,
)
from app.services.procurement_pricing_source import build_table, current_prices, price_history

router = APIRouter(prefix="/procurement-order-formation/pricing", tags=["procurement-pricing"])


def buyer(
    session: ProcurementOrderFormationSession = Depends(verify_procurement_order_formation_session),
):
    ensure_bitrix_user_allowed(session.user_id)
    return session


def failure(exc):
    if isinstance(exc, PricingConflict):
        return HTTPException(409, str(exc))
    if isinstance(exc, LookupError):
        return HTTPException(404, str(exc))
    if isinstance(exc, ValueError):
        return HTTPException(422, str(exc))
    return HTTPException(503, "Источник данных временно недоступен; изменения не применены")


@router.get("/table", response_model=PricingTable)
def table(
    filters: PricingFilter = Depends(), db: Session = Depends(get_db), session=Depends(buyer)
):
    try:
        return build_table(db, filters)
    except Exception as exc:
        raise failure(exc) from exc


@router.get("/history/{code}", response_model=list[PriceHistoryPoint])
def history(code: str, start: date, end: date, session=Depends(buyer)):
    try:
        PricingFilter(start=start, end=end)
        return price_history(code, start, end)
    except Exception as exc:
        raise failure(exc) from exc


def safe_cell(value):
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + value
    return value


@router.get("/export")
def export(
    format: Literal["csv", "xlsx"] = "xlsx",
    filters: PricingFilter = Depends(),
    db: Session = Depends(get_db),
    session=Depends(buyer),
):
    try:
        result = build_table(db, filters, paginate=False)
    except Exception as exc:
        raise failure(exc) from exc
    fields = (
        list(result.items[0].model_fields)
        if result.items
        else [
            "code",
            "name",
            "bronze",
            "platinum",
            "sales_qty",
            "sales_amount",
            "forecast_qty",
            "forecast_amount",
        ]
    )
    rows = [
        ["Период", str(result.start), str(result.end)],
        ["Обновлено", result.observed_at.isoformat()],
        fields,
    ]
    for item in result.items:
        values = item.model_dump(mode="json")
        rows.append(
            [safe_cell(str(values[key])) if values[key] is not None else "" for key in fields]
        )
    if format == "csv":
        stream = StringIO()
        csv.writer(stream, delimiter=";").writerows(rows)
        data = stream.getvalue().encode("utf-8-sig")
        mime = "text/csv; charset=utf-8"
    else:
        book = Workbook()
        sheet = book.active
        sheet.title = "Ценообразование"
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes = "A4"
        sheet.auto_filter.ref = f"A3:{sheet.cell(3,len(fields)).column_letter}{sheet.max_row}"
        stream = BytesIO()
        book.save(stream)
        data = stream.getvalue()
        mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return Response(
        data,
        media_type=mime,
        headers={
            "Content-Disposition": f'attachment; filename="pricing-{result.start}-{result.end}.{format}"'
        },
    )


@router.get("/batches", response_model=list[PriceBatchRead])
def batches(db: Session = Depends(get_db), session=Depends(buyer)):
    return db.scalars(
        select(ProcurementPriceBatch)
        .where(ProcurementPriceBatch.owner == actor_key(session))
        .order_by(ProcurementPriceBatch.id.desc())
        .limit(100)
    ).all()


@router.post("/batches", response_model=PriceBatchRead)
def create(payload: PriceBatchCreate, db: Session = Depends(get_db), session=Depends(buyer)):
    try:
        batch = create_batch(db, payload, actor_key(session))
        db.commit()
        return batch
    except Exception as exc:
        db.rollback()
        raise failure(exc) from exc


@router.get("/batches/{batch_id}", response_model=PriceBatchRead)
def read(batch_id: int, db: Session = Depends(get_db), session=Depends(buyer)):
    try:
        return get_batch(db, batch_id, actor_key(session))
    except Exception as exc:
        raise failure(exc) from exc


@router.post("/batches/{batch_id}/approve", response_model=PriceBatchRead)
def approve(
    batch_id: int, payload: PriceApproval, db: Session = Depends(get_db), session=Depends(buyer)
):
    try:
        result = approve_batch(db, batch_id, payload.version, actor_key(session), current_prices)
        db.commit()
        return result
    except Exception as exc:
        db.rollback()
        raise failure(exc) from exc


@router.get("/presets", response_model=list[PricePresetRead])
def presets(db: Session = Depends(get_db), session=Depends(buyer)):
    return [
        {"id": x.id, "name": x.name, "filters": x.filters}
        for x in db.scalars(
            select(ProcurementPricePreset)
            .where(ProcurementPricePreset.owner == actor_key(session))
            .order_by(ProcurementPricePreset.name)
        ).all()
    ]


@router.post("/presets", response_model=PricePresetRead)
def save_preset(payload: PricePresetWrite, db: Session = Depends(get_db), session=Depends(buyer)):
    preset = db.scalar(
        select(ProcurementPricePreset).where(
            ProcurementPricePreset.owner == actor_key(session),
            ProcurementPricePreset.name == payload.name,
        )
    )
    if preset is None:
        preset = ProcurementPricePreset(owner=actor_key(session), name=payload.name)
        db.add(preset)
    preset.filters = payload.filters.model_dump(mode="json")
    db.commit()
    return {"id": preset.id, "name": preset.name, "filters": preset.filters}


@router.delete("/presets/{preset_id}", status_code=204)
def delete_preset(preset_id: int, db: Session = Depends(get_db), session=Depends(buyer)):
    preset = db.scalar(
        select(ProcurementPricePreset).where(
            ProcurementPricePreset.id == preset_id,
            ProcurementPricePreset.owner == actor_key(session),
        )
    )
    if preset is None:
        raise HTTPException(404, "Фильтр не найден")
    db.delete(preset)
    db.commit()
